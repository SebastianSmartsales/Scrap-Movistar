"""
Arma data/rucs_movistar.xlsx con lo que haya en data/.

Hoja "Resumen": una fila por RUC con RUC | Razón social | Tipo de cliente | Subtipo de cliente y, si existe
data/productos.csv, un bloque de columnas por subtipo de producto (MOVISTAR, MONO, DUO, TRIO) bajo un encabezado
combinado: Cant. | Códigos (uno por línea) | Planes (distintos, con su cantidad).
Hoja "Detalle": una fila por producto (RUC | Razón social | Subtipo de producto | Código de producto | Plan).
"""
import json
import os

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import sesion

CSV = os.path.join(sesion.DATA, "rucs.csv")
XLSX = os.path.join(sesion.DATA, "rucs_movistar.xlsx")
PROGRESO = os.path.join(sesion.DATA, "progreso.json")
PRODUCTOS = os.path.join(sesion.DATA, "productos.csv")

BASE = ["RUC", "Razón social", "Tipo de cliente", "Subtipo de cliente"]
ANCHOS_BASE = [14, 60, 20, 30]
SUBTIPOS = ["MOVISTAR", "MONO", "DUO", "TRIO"]
COLORES = {"MOVISTAR": "019DF4", "MONO": "5CB615", "DUO": "A64DB5", "TRIO": "F08A00"}  # azul, verde, morado, naranja
MAX_CELDA = 30_000  # Excel admite 32.767 caracteres por celda
POR_HOJA = 1_000_000  # el límite de Excel es 1.048.576 filas

AJUSTE = Alignment(wrap_text=True, vertical="top")
CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)


def leer() -> pd.DataFrame:
    """Une todas las pasadas (data/pasada*.csv y data/rucs.csv) y limpia los BOM que dejó la primera versión."""
    archivos = sorted(f for f in os.listdir(sesion.DATA) if f.startswith("pasada") and f.endswith(".csv"))
    rutas = [os.path.join(sesion.DATA, f) for f in archivos] + [CSV]
    print("Pasadas:", [os.path.basename(r) for r in rutas])
    df = pd.concat([pd.read_csv(r, dtype=str, keep_default_na=False, encoding="utf-8-sig") for r in rutas])
    return df.apply(lambda c: c.str.replace("﻿", "", regex=False).str.strip())


def recortar(lineas: list) -> str:
    """Une las líneas con salto de línea; si no caben en una celda, corta y avisa que el resto está en Detalle."""
    texto, usado = [], 0
    for i, linea in enumerate(lineas):
        if usado + len(linea) + 1 > MAX_CELDA:
            texto.append(f"… (+{len(lineas) - i} más, ver hoja Detalle)")
            break
        texto.append(linea)
        usado += len(linea) + 1
    return "\n".join(texto)


def bloques(prod: pd.DataFrame) -> dict:
    """{ruc: {subtipo: (cantidad, códigos, planes)}} a partir de las filas de producto."""
    out = {}
    for (ruc, sub), g in prod.groupby(["ruc", "subtipo_producto"], sort=False):
        planes = {}
        for plan in g["plan"]:
            planes[plan] = planes.get(plan, 0) + 1
        lineas_plan = [f"{p} (x{n})" if n > 1 else p for p, n in planes.items()]
        out.setdefault(ruc, {})[sub] = (len(g), recortar(list(g["codigo_producto"])), recortar(lineas_plan))
    return out


def celda(ws, valor, estilo=None, fuente=None, relleno=None):
    c = WriteOnlyCell(ws, value=valor)
    if estilo:
        c.alignment = estilo
    if fuente:
        c.font = fuente
    if relleno:
        c.fill = relleno
    return c


def hoja_resumen(wb, out: pd.DataFrame, por_ruc) -> None:
    ws = wb.create_sheet("Resumen")
    blanco = Font(bold=True, color="FFFFFF")
    gris = PatternFill("solid", fgColor="404040")
    subtipos = SUBTIPOS if por_ruc is not None else []
    ancho = list(ANCHOS_BASE)
    for _ in subtipos:
        ancho += [7, 18, 42]
    if subtipos:
        ancho += [10]  # total de productos
    for i, a in enumerate(ancho, 1):
        ws.column_dimensions[get_column_letter(i)].width = a
    # encabezado en dos filas: fila 1 = bandas por subtipo, fila 2 = nombres de columna
    fila1 = [celda(ws, n, CENTRO, blanco, gris) for n in BASE]
    fila2 = [celda(ws, None, CENTRO, blanco, gris) for _ in BASE]
    for sub in subtipos:
        relleno = PatternFill("solid", fgColor=COLORES[sub])
        fila1 += [celda(ws, sub, CENTRO, blanco, relleno), celda(ws, None, CENTRO, blanco, relleno),
                  celda(ws, None, CENTRO, blanco, relleno)]
        fila2 += [celda(ws, n, CENTRO, blanco, relleno) for n in ("Cant.", "Códigos", "Planes")]
    if subtipos:
        fila1.append(celda(ws, "Total", CENTRO, blanco, gris))
        fila2.append(celda(ws, "productos", CENTRO, blanco, gris))
    ws.freeze_panes = "C3"  # en modo escritura rápida va antes de la primera fila
    ws.append(fila1)
    ws.append(fila2)
    for i in range(1, len(BASE) + 1):
        ws.merged_cells.add(f"{get_column_letter(i)}1:{get_column_letter(i)}2")
    for k in range(len(subtipos)):
        c0 = len(BASE) + 1 + 3 * k
        ws.merged_cells.add(f"{get_column_letter(c0)}1:{get_column_letter(c0 + 2)}1")
    if subtipos:
        c = len(BASE) + 1 + 3 * len(subtipos)
        ws.merged_cells.add(f"{get_column_letter(c)}1:{get_column_letter(c)}2")
    for ruc, razon, tipo, subt in out.itertuples(index=False):
        fila = [ruc, razon, tipo, subt]
        if por_ruc is not None:
            total = 0
            datos = por_ruc.get(ruc, {})
            for sub in subtipos:
                n, codigos, planes = datos.get(sub, (None, "", ""))
                total += n or 0
                fila += [n, celda(ws, codigos, AJUSTE) if codigos else None, celda(ws, planes, AJUSTE) if planes else None]
            fila.append(total or None)
        ws.append(fila)
    ws.auto_filter.ref = f"A2:{get_column_letter(len(ancho))}{len(out) + 2}"


def hoja_detalle(wb, out: pd.DataFrame, prod: pd.DataFrame) -> None:
    det = prod.merge(out[["ruc", "razon_social"]], on="ruc", how="left")
    det = det[["ruc", "razon_social", "subtipo_producto", "codigo_producto", "plan"]]
    cols = ["RUC", "Razón social", "Subtipo de producto", "Código de producto", "Plan"]
    anchos = [14, 60, 20, 18, 45]
    for n, i in enumerate(range(0, max(len(det), 1), POR_HOJA), 1):
        ws = wb.create_sheet("Detalle" if len(det) <= POR_HOJA else f"Detalle {n}")
        for j, a in enumerate(anchos, 1):
            ws.column_dimensions[get_column_letter(j)].width = a
        ws.append([celda(ws, c, CENTRO, Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="404040"))
                   for c in cols])
        ws.freeze_panes = "A2"
        trozo = det.iloc[i:i + POR_HOJA]
        for fila in trozo.itertuples(index=False):
            ws.append(list(fila))
        ws.auto_filter.ref = f"A1:E{len(trozo) + 1}"


def main() -> None:
    df = leer()
    leidas = len(df)
    df = df.drop_duplicates(subset="ruc", keep="first")
    print(f"Filas leídas: {leidas}; únicas por RUC: {len(df)}; duplicadas: {leidas - len(df)}")
    if os.path.exists(PROGRESO):
        total = json.load(open(PROGRESO)).get("total_portal")
        print(f"Total que informó el portal: {total}; diferencia: {total - len(df) if total else '?'}")
    df = df[df["tipo_doc"] == "RUC"]  # fuera los registros de prueba (tipo_doc OTROS)
    print(f"Solo RUC reales: {len(df)}")
    print("Por estado:", df["estado"].value_counts().to_dict())
    print("Por tipo de cliente:", df["tipo_cliente"].replace("", "(vacío)").value_counts().to_dict())

    out = df[["ruc", "razon_social", "tipo_cliente", "subtipo_cliente"]]
    prod = por_ruc = None
    if os.path.exists(PRODUCTOS):
        todo = pd.read_csv(PRODUCTOS, dtype=str, keep_default_na=False, encoding="utf-8-sig")
        prod = todo[todo["estado"] == "ok"][["ruc", "subtipo_producto", "codigo_producto", "plan"]]
        print(f"RUC con productos consultados: {todo['ruc'].nunique()} de {len(out)}; "
              f"RUC con productos activos: {prod['ruc'].nunique()}; filas de producto: {len(prod)}")
        print("Por subtipo de producto:", prod["subtipo_producto"].value_counts().to_dict())
        por_ruc = bloques(prod)

    wb = Workbook(write_only=True)
    hoja_resumen(wb, out, por_ruc)
    if prod is not None:
        hoja_detalle(wb, out, prod)
    wb.save(XLSX)
    print(f"Guardado: {XLSX}")


if __name__ == "__main__":
    main()
